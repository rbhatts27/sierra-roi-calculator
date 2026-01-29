from flask import Flask, request, jsonify, send_from_directory
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

app = Flask(__name__, static_folder='static')

EXCEL_FILE = 'roi_data.xlsx'

def init_excel():
    """Initialize Excel file with headers if it doesn't exist."""
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "ROI Data"
        headers = [
            'Timestamp', 'Customer Name', 'Account SID',
            'Current Annual Cost', 'Expected Cost Reduction (%)',
            'Implementation Cost', 'Annual Maintenance Cost',
            'Time to Value (months)', 'Calculated ROI (%)',
            'Net Savings', 'Payback Period (months)', 'Notes'
        ]
        ws.append(headers)
        wb.save(EXCEL_FILE)

def calculate_roi(data):
    """Calculate ROI metrics."""
    current_cost = float(data.get('current_annual_cost', 0))
    cost_reduction_pct = float(data.get('expected_cost_reduction', 0)) / 100
    implementation_cost = float(data.get('implementation_cost', 0))
    annual_maintenance = float(data.get('annual_maintenance_cost', 0))

    annual_savings = current_cost * cost_reduction_pct
    net_annual_savings = annual_savings - annual_maintenance
    total_first_year_cost = implementation_cost + annual_maintenance

    # ROI calculation (first year)
    if total_first_year_cost > 0:
        roi = ((net_annual_savings - implementation_cost) / total_first_year_cost) * 100
    else:
        roi = 0

    # Payback period in months
    if net_annual_savings > 0:
        payback_months = (implementation_cost / net_annual_savings) * 12
    else:
        payback_months = 0

    return {
        'roi_percentage': round(roi, 2),
        'net_savings': round(net_annual_savings, 2),
        'payback_months': round(payback_months, 1)
    }

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')

@app.route('/api/submit', methods=['POST'])
def submit_roi():
    """Save ROI data to Excel."""
    init_excel()
    data = request.json

    # Calculate ROI
    roi_results = calculate_roi(data)

    # Load and append to Excel
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    row = [
        datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        data.get('customer_name', ''),
        data.get('account_sid', ''),
        float(data.get('current_annual_cost', 0)),
        float(data.get('expected_cost_reduction', 0)),
        float(data.get('implementation_cost', 0)),
        float(data.get('annual_maintenance_cost', 0)),
        float(data.get('time_to_value', 0)),
        roi_results['roi_percentage'],
        roi_results['net_savings'],
        roi_results['payback_months'],
        data.get('notes', '')
    ]
    ws.append(row)
    wb.save(EXCEL_FILE)

    return jsonify({
        'success': True,
        'message': 'Data saved successfully',
        'roi_results': roi_results
    })

@app.route('/api/search', methods=['GET'])
def search_roi():
    """Search for ROI data by customer name or account SID."""
    init_excel()

    customer_name = request.args.get('customer_name', '').lower().strip()
    account_sid = request.args.get('account_sid', '').lower().strip()

    if not customer_name and not account_sid:
        return jsonify({'success': False, 'message': 'Please provide customer name or account SID'})

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    results = []
    headers = [cell.value for cell in ws[1]]

    for row in ws.iter_rows(min_row=2, values_only=True):
        row_customer = str(row[1]).lower() if row[1] else ''
        row_sid = str(row[2]).lower() if row[2] else ''

        if (customer_name and customer_name in row_customer) or \
           (account_sid and account_sid in row_sid):
            results.append(dict(zip(headers, row)))

    return jsonify({
        'success': True,
        'count': len(results),
        'results': results
    })

@app.route('/api/all', methods=['GET'])
def get_all():
    """Get all ROI entries."""
    init_excel()

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    results = []
    headers = [cell.value for cell in ws[1]]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(row):  # Skip empty rows
            results.append(dict(zip(headers, row)))

    return jsonify({
        'success': True,
        'count': len(results),
        'results': results
    })

if __name__ == '__main__':
    init_excel()
    print("Starting Sierra ROI Calculator...")
    print("Open http://localhost:5000 in your browser")
    app.run(debug=True, port=5000)
